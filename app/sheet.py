"""Read the expenses table from Google Drive (.xlsx file or native Google Sheet) and parse rows."""
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import openpyxl
from google.auth.transport.requests import AuthorizedSession
from google.oauth2 import service_account

DRIVE = "https://www.googleapis.com/drive/v3"
SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GSHEET_MIME = "application/vnd.google-apps.spreadsheet"

# logical column → accepted header names (lower-case)
DEFAULT_HEADERS = {
    "number": ["№", "no", "nr", "number", "expenses №", "номер"],
    "date": ["date", "дата", "tanggal"],
    "payee": ["payee", "vendor", "supplier", "поставщик", "получатель"],
    "amount": ["amount", "сумма", "jumlah"],
    "currency": ["currency", "валюта"],
    "type": ["type", "тип"],
    "loan": ["paid from director personal account (loan)", "loan", "заём"],
    "description": ["description", "описание"],
    "main_doc": ["main document", "document", "документ"],
    "supporting_docs": ["supporting documents"],
    "drive_folder": ["drive folder", "папка"],
    "notes": ["notes", "примечание", "заметки"],
}
REQUIRED = ["number", "date", "payee", "amount", "currency"]


@dataclass
class Row:
    number: int
    sheet_row: int
    date: date | None
    payee: str
    amount: float | None
    currency: str
    type: str = ""
    loan: str = ""
    description: str = ""
    main_doc: str = ""
    supporting_docs: list[str] = field(default_factory=list)
    folder_id: str = ""
    notes: str = ""
    errors: list[str] = field(default_factory=list)

    def summary(self) -> dict:
        return {"no": self.number, "date": self.date.isoformat() if self.date else None, "payee": self.payee,
                "amount": self.amount, "currency": self.currency, "type": self.type,
                "description": self.description[:120]}


class Drive:
    def __init__(self, sa_json: str):
        info = json.loads(sa_json)
        creds = service_account.Credentials.from_service_account_info(info, scopes=SCOPES)
        self.email = info.get("client_email", "")
        self.s = AuthorizedSession(creds)

    def _get(self, url, **params):
        params.setdefault("supportsAllDrives", "true")
        r = self.s.get(url, params=params, timeout=60)
        if r.status_code == 404:
            raise FileNotFoundError(
                f"Not found or no access. Share the file/folder with the service account: {self.email}")
        r.raise_for_status()
        return r

    def meta(self, file_id: str) -> dict:
        return self._get(f"{DRIVE}/files/{file_id}", fields="id,name,mimeType,modifiedTime").json()

    def download_table(self, file_id: str) -> tuple[bytes, dict]:
        m = self.meta(file_id)
        if m["mimeType"] == GSHEET_MIME:
            data = self._get(f"{DRIVE}/files/{file_id}/export", mimeType=XLSX_MIME).content
        else:
            data = self._get(f"{DRIVE}/files/{file_id}", alt="media").content
        return data, m

    def find_in_folder(self, folder_id: str, name: str) -> dict | None:
        safe = name.replace("\\", "\\\\").replace("'", "\\'")
        q = f"'{folder_id}' in parents and name = '{safe}' and trashed = false"
        files = self._get(f"{DRIVE}/files", q=q, fields="files(id,name,mimeType,size)",
                          includeItemsFromAllDrives="true", pageSize=5).json().get("files", [])
        return files[0] if files else None

    def download(self, file_id: str) -> bytes:
        return self._get(f"{DRIVE}/files/{file_id}", alt="media").content


# ---------------------------------------------------------------- parsing
def parse_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):  # Excel serial
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
    s = str(v).strip()
    for cand in (s, s[:10]):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y"):
            try:
                return datetime.strptime(cand, fmt).date()
            except ValueError:
                continue
    try:
        return datetime.fromisoformat(s).date()
    except ValueError:
        return None


def parse_amount(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not s:
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        th = "." if dec == "," else ","
        s = s.replace(th, "").replace(dec, ".")
    elif s.count(",") > 1 or s.count(".") > 1:
        s = s.replace(",", "").replace(".", "")
    elif "," in s or "." in s:
        sep = "," if "," in s else "."
        tail = s.split(sep)[1]
        s = s.replace(sep, "") if len(tail) == 3 else s.replace(sep, ".")
    try:
        return float(s)
    except ValueError:
        return None


FOLDER_RE = re.compile(r"/folders/([A-Za-z0-9_\-]+)")


def _folder_id(formula, value) -> str:
    for x in (formula, value):
        if isinstance(x, str):
            m = FOLDER_RE.search(x)
            if m:
                return m.group(1)
    return ""


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def parse_table(data: bytes, sheet_name: str, column_map: dict | None = None) -> list[Row]:
    wb_v = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    wb_f = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    name = sheet_name if sheet_name in wb_v.sheetnames else wb_v.sheetnames[0]
    vals = list(wb_v[name].iter_rows(values_only=True))
    forms = list(wb_f[name].iter_rows(values_only=True))
    if not vals:
        return []

    headers = [(_s(h)).lower() for h in vals[0]]
    col: dict[str, int] = {}
    wanted = {k: list(v) for k, v in DEFAULT_HEADERS.items()}
    for k, h in (column_map or {}).items():
        wanted[k] = [h.lower()]
    for key, names in wanted.items():
        for i, h in enumerate(headers):
            if h in names:
                col[key] = i
                break
    miss = [k for k in REQUIRED if k not in col]
    if miss:
        raise ValueError(f"Columns not found in sheet '{name}': {miss}. Headers: {headers}. Use COLUMN_MAP.")

    def g(r, key):
        i = col.get(key)
        return r[i] if i is not None and i < len(r) else None

    rows: list[Row] = []
    for idx, r in enumerate(vals[1:], start=2):
        num = g(r, "number")
        if num in (None, "") or not str(num).strip().replace(".0", "").isdigit():
            continue
        f = forms[idx - 1] if idx - 1 < len(forms) else ()
        fi = col.get("drive_folder")
        row = Row(
            number=int(float(num)), sheet_row=idx,
            date=parse_date(g(r, "date")), payee=_s(g(r, "payee")),
            amount=parse_amount(g(r, "amount")), currency=_s(g(r, "currency")).upper(),
            type=_s(g(r, "type")), loan=_s(g(r, "loan")), description=_s(g(r, "description")),
            main_doc=_s(g(r, "main_doc")),
            supporting_docs=[x.strip() for x in _s(g(r, "supporting_docs")).split(";") if x.strip()],
            folder_id=_folder_id(f[fi] if fi is not None and fi < len(f) else None, g(r, "drive_folder")),
            notes=_s(g(r, "notes")),
        )
        if not row.date:
            row.errors.append("no/invalid date")
        if not row.payee:
            row.errors.append("no payee")
        if row.amount is None or row.amount <= 0:
            row.errors.append("no/invalid amount")
        if not re.fullmatch(r"[A-Z]{3}", row.currency):
            row.errors.append("no/invalid currency")
        rows.append(row)
    return rows
